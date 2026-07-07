"""Tests for document file-content analysis (file_processor + file_schemas)."""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest

from src.llm.file_schemas import (
    FILE_RISK_TOOL_NAME,
    FileCategory,
    FileRiskAnalysis,
    FileRiskFinding,
    build_file_risk_tool,
)
from src.pipeline.file_processor import _extract_text, _is_benign_partner_report

# ---------------------------------------------------------------------------
# Schema tests
# ---------------------------------------------------------------------------


def test_file_risk_finding_parses_correctly() -> None:
    data = {
        "category": "credentials",
        "excerpt": "password=secret123",
        "explanation": "Plaintext password exposed.",
        "score": 90,
        "confidence": 0.95,
    }
    f = FileRiskFinding.model_validate(data)
    assert f.category == FileCategory.CREDENTIALS
    assert f.score == 90
    assert f.confidence == 0.95


def test_file_risk_analysis_empty_by_default() -> None:
    a = FileRiskAnalysis()
    assert a.findings == []


def test_file_risk_analysis_parses_multiple_findings() -> None:
    data = {
        "findings": [
            {
                "category": "financial",
                "excerpt": "revenue: $5M",
                "explanation": "Exact revenue figure.",
                "score": 70,
                "confidence": 0.8,
            },
            {
                "category": "credentials",
                "excerpt": "api_key=abc123",
                "explanation": "API key exposed.",
                "score": 95,
                "confidence": 0.99,
            },
        ]
    }
    a = FileRiskAnalysis.model_validate(data)
    assert len(a.findings) == 2
    assert a.findings[0].category == FileCategory.FINANCIAL
    assert a.findings[1].category == FileCategory.CREDENTIALS


def test_file_category_enum_values() -> None:
    assert FileCategory.CREDENTIALS == "credentials"
    assert FileCategory.FINANCIAL == "financial"
    assert FileCategory.PERSONAL_DATA == "personal_data"
    assert FileCategory.BUSINESS_SECRETS == "business_secrets"
    assert FileCategory.INTERNAL_INFRA == "internal_infra"
    assert FileCategory.LEGAL == "legal"


def test_build_file_risk_tool_structure() -> None:
    tool = build_file_risk_tool()
    assert tool["type"] == "function"
    assert tool["function"]["name"] == FILE_RISK_TOOL_NAME
    assert "parameters" in tool["function"]
    schema = tool["function"]["parameters"]
    assert "findings" in schema["properties"]


def test_file_risk_finding_rejects_extra_fields() -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        FileRiskFinding.model_validate(
            {
                "category": "credentials",
                "excerpt": "x",
                "explanation": "y",
                "score": 50,
                "confidence": 0.5,
                "unexpected_field": True,
            }
        )


def test_file_risk_finding_score_bounds() -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        FileRiskFinding.model_validate(
            {
                "category": "financial",
                "excerpt": "x",
                "explanation": "y",
                "score": 150,
                "confidence": 0.5,
            }
        )


# ---------------------------------------------------------------------------
# Text extraction tests
# ---------------------------------------------------------------------------


def test_extract_text_plain_txt() -> None:
    data = b"hello world\nline two"
    result = _extract_text(data, "report.txt", "text/plain")
    assert result == "hello world\nline two"


def test_extract_text_csv() -> None:
    data = b"name,value\nfoo,bar"
    result = _extract_text(data, "data.csv", "")
    assert result is not None
    assert "foo" in result


def test_extract_text_json() -> None:
    data = b'{"key": "value"}'
    result = _extract_text(data, "config.json", "application/json")
    assert result is not None
    assert "value" in result


def test_extract_text_docx() -> None:
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("Secret salary: $200k")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    result = _extract_text(
        data,
        "report.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert result is not None
    assert "Secret salary" in result


def test_extract_text_pdf() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    data = buf.getvalue()

    # Blank PDF: extraction succeeds but is empty (that's fine — skipped upstream)
    result = _extract_text(data, "doc.pdf", "application/pdf")
    assert result is not None  # no crash; empty is handled by caller


def test_extract_text_xlsx() -> None:
    import openpyxl  # type: ignore[import-untyped]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Partner", "Revenue"])  # type: ignore[union-attr]
    ws.append(["ACME", "5000000"])  # type: ignore[union-attr]
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    result = _extract_text(
        data,
        "sheet.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert result is not None
    assert "ACME" in result
    assert "5000000" in result


def test_extract_text_unsupported_returns_none() -> None:
    result = _extract_text(b"\x89PNG\r\n", "image.png", "image/png")
    assert result is None


def test_extract_text_truncation_handled_by_caller() -> None:
    # _extract_text returns raw text; truncation is done in _analyze_file
    data = ("x" * 50_000).encode()
    result = _extract_text(data, "big.txt", "text/plain")
    assert result is not None
    assert len(result) == 50_000


# ---------------------------------------------------------------------------
# benign partner report suppression (2026-07-06 confirmed FP)
# ---------------------------------------------------------------------------


def test_benign_partner_report_matches_commission_and_payout() -> None:
    # The confirmed-FP artifact and its family are suppressed by filename.
    assert _is_benign_partner_report("player_commission_report_06-07-2026.xlsx")
    assert _is_benign_partner_report("Commission Report Q2.xlsx")
    assert _is_benign_partner_report("payout-report.csv")
    assert _is_benign_partner_report("Payout_Report.pdf")


def test_benign_partner_report_does_not_over_match() -> None:
    # Precise to commission/payout reports — every other document (including other
    # data_leak artifacts) is still analysed and can alert.
    assert not _is_benign_partner_report("credentials.txt")
    assert not _is_benign_partner_report("internal_strategy.docx")
    assert not _is_benign_partner_report("player_list.xlsx")
    assert not _is_benign_partner_report("report_2026.pdf")   # bare 'report'
    assert not _is_benign_partner_report("commissions.xlsx")  # needs 'report' too


@pytest.mark.asyncio
async def test_analyze_file_skips_benign_report_before_spending(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # A commission-report document is skipped BEFORE download/extraction/LLM, so it
    # never becomes a risk_event or an alert (and costs nothing).
    monkeypatch.setattr("src.pipeline.file_processor.settings.FILE_ANALYSIS_ENABLED", True)

    complete_mock = AsyncMock()
    monkeypatch.setattr("src.pipeline.file_processor.complete_task", complete_mock)

    msg = MagicMock()
    msg.id = uuid4()
    msg.chat_id = uuid4()
    msg.sender_id = 111
    msg.raw_payload = {
        "document": {
            "file_id": "F1",
            "file_name": "player_commission_report_06-07-2026.xlsx",
            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }
    }
    monkeypatch.setattr(
        "src.pipeline.file_processor.get_message_by_id", AsyncMock(return_value=msg)
    )
    monkeypatch.setattr(
        "src.pipeline.file_processor.get_chat_by_id", AsyncMock(return_value=MagicMock())
    )

    download_mock = AsyncMock()
    analyze_mock = AsyncMock()
    monkeypatch.setattr("src.pipeline.file_processor._download", download_mock)
    monkeypatch.setattr("src.pipeline.file_processor.analyze_file_risk", analyze_mock)

    fake_conn = AsyncMock()
    fake_ctx = AsyncMock()
    fake_ctx.__aenter__ = AsyncMock(return_value=fake_conn)
    fake_ctx.__aexit__ = AsyncMock(return_value=False)
    monkeypatch.setattr(
        "src.pipeline.file_processor.acquire_connection", MagicMock(return_value=fake_ctx)
    )

    from src.db.models import ProcessingQueue
    from src.pipeline.file_processor import process_file_task

    task = MagicMock(spec=ProcessingQueue)
    task.id = 7
    task.payload = {"message_id": str(msg.id)}
    task.attempts = 1

    await process_file_task(MagicMock(), task)

    complete_mock.assert_awaited_once()      # task completed (skip path)
    download_mock.assert_not_called()        # never downloaded
    analyze_mock.assert_not_called()         # never sent to the LLM


# ---------------------------------------------------------------------------
# process_file_task — kill-switch test
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_process_file_task_disabled(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("src.pipeline.file_processor.settings.FILE_ANALYSIS_ENABLED", False)

    complete_mock = AsyncMock()
    monkeypatch.setattr("src.pipeline.file_processor.complete_task", complete_mock)

    fake_conn = AsyncMock()
    fake_ctx = AsyncMock()
    fake_ctx.__aenter__ = AsyncMock(return_value=fake_conn)
    fake_ctx.__aexit__ = AsyncMock(return_value=False)
    monkeypatch.setattr(
        "src.pipeline.file_processor.acquire_connection", MagicMock(return_value=fake_ctx)
    )

    from src.db.models import ProcessingQueue
    from src.pipeline.file_processor import process_file_task

    task = MagicMock(spec=ProcessingQueue)
    task.id = 1
    task.payload = {"message_id": str(uuid4())}
    task.attempts = 1

    bot = MagicMock()
    await process_file_task(bot, task)

    complete_mock.assert_awaited_once()


@pytest.mark.asyncio
async def test_process_file_task_missing_message_id(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("src.pipeline.file_processor.settings.FILE_ANALYSIS_ENABLED", True)

    fail_mock = AsyncMock()
    monkeypatch.setattr("src.pipeline.file_processor.fail_task", fail_mock)

    fake_conn = AsyncMock()
    fake_ctx = AsyncMock()
    fake_ctx.__aenter__ = AsyncMock(return_value=fake_conn)
    fake_ctx.__aexit__ = AsyncMock(return_value=False)
    monkeypatch.setattr(
        "src.pipeline.file_processor.acquire_connection", MagicMock(return_value=fake_ctx)
    )

    from src.db.models import ProcessingQueue
    from src.pipeline.file_processor import process_file_task

    task = MagicMock(spec=ProcessingQueue)
    task.id = 2
    task.payload = {}  # no message_id
    task.attempts = 1

    bot = MagicMock()
    await process_file_task(bot, task)

    fail_mock.assert_awaited_once()
